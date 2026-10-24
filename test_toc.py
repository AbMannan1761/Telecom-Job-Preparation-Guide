import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = docx.Document()
tbl = doc.add_table(rows=1, cols=2)
p1 = tbl.cell(0, 0).paragraphs[0]
p2 = tbl.cell(0, 1).paragraphs[0]

hl = parse_xml(f'<w:hyperlink {nsdecls("w")} w:anchor="test_bm" w:history="1"><w:r><w:t>Chapter 1</w:t></w:r></w:hyperlink>')
p1._p.append(hl)

fld = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGEREF test_bm \\h"><w:r><w:t>5</w:t></w:r></w:fldSimple>')
p2._p.append(fld)

doc.add_page_break()
p_target = doc.add_paragraph()
bm_start = parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="10" w:name="test_bm"/>')
bm_end = parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="10"/>')
p_target._p.append(bm_start)
p_target.add_run('This is target')
p_target._p.append(bm_end)

doc.save('test_toc_table.docx')
print('XML CREATION SUCCESS!')
